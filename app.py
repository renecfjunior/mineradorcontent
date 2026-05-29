# -*- coding: utf-8 -*-
"""
ELEMENTAR — Minerador de Pautas
Coleta notícias do G1, Nexo e Agência IBGE, pontua cada uma pela Matriz de
Validação do dossiê e sugere um ângulo de vídeo. Exporta para Word (.docx).

Rode com:  python app.py   →   http://127.0.0.1:5000
"""

import re
import io
import html
from datetime import datetime

import requests
import feedparser
from flask import Flask, render_template, jsonify, request, send_file

from docx import Document
from docx.shared import Pt, RGBColor

app = Flask(__name__, template_folder=".")

# ---------------------------------------------------------------------------
# 1. FONTES (cada fonte tem candidatos de feed; usa o primeiro que responder)
# ---------------------------------------------------------------------------
FONTES = {
    "G1": [
        "https://g1.globo.com/rss/g1/economia/",
        "https://g1.globo.com/rss/g1/politica/",
        "https://g1.globo.com/rss/g1/brasil/",
        "http://g1.globo.com/dynamo/economia/rss2.xml",
    ],
    "Nexo": [
        "https://www.nexojornal.com.br/feed/",
        "https://www.nexojornal.com.br/rss",
        "https://www.nexojornal.com.br/feed.rss",
    ],
    "IBGE": [
        "https://agenciadenoticias.ibge.gov.br/agencia-noticias/2012-agencia-de-noticias/noticias?format=feed&type=rss",
        "https://agenciadenoticias.ibge.gov.br/component/obrss/rss-agencia",
    ],
}

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                  "AppleWebKit/537.36 (KHTML, like Gecko) "
                  "Chrome/124.0 Safari/537.36"
}

# ---------------------------------------------------------------------------
# 2. DICIONÁRIOS DE PONTUAÇÃO  (edite à vontade)
# ---------------------------------------------------------------------------

# Categorias temáticas do dossiê (seção B)
CATEGORIAS = {
    "Monopólios e Poder Econômico": [
        "cade", "antitruste", "monopólio", "monopolio", "cartel", "fusão", "fusao",
        "aquisição", "aquisicao", "fraude fiscal", "abuso de poder", "concentração",
        "concentracao", "domínio", "dominio", "oligopólio", "oligopolio",
        "conluio", "bilionária", "bilionaria", "truste",
    ],
    "Infraestrutura e Engenharia": [
        "tcu", "obra", "obras", "pac", "abnt", "incc", "infraestrutura",
        "licitação", "licitacao", "concessão", "concessao", "ferrovia",
        "rodovia", "saneamento", "porto", "superfaturamento", "metrô", "metro",
        "construção", "construcao",
    ],
    "Indústria do Consumo e Margens": [
        "margem de lucro", "custo de produção", "custo de producao", "falsificação",
        "falsificacao", "pirataria", "obsolescência", "obsolescencia", "apreensão",
        "apreensao", "contrabando", "markup", "lucro abusivo", "cosméticos",
        "cosmeticos", "eletrônicos", "eletronicos", "recall",
    ],
    "Declínio Urbano e Imobiliário": [
        "imóveis vagos", "imoveis vagos", "censo", "crise habitacional",
        "êxodo", "exodo", "gentrificação", "gentrificacao", "vacância", "vacancia",
        "esvaziamento", "imobiliário", "imobiliario", "cracolândia", "cracolandia",
        "déficit habitacional", "deficit habitacional", "favela",
    ],
}

# Critério 1 — Poder de Gancho: conflito + indignação + curiosidade
GANCHO = [
    "monopólio", "monopolio", "escândalo", "escandalo", "esquema", "secreto",
    "oculto", "bilionária", "bilionaria", "ninguém", "ninguem", "por que",
    "recorde", "verdade", "bastidores", "cartel", "fraude", "golpe", "colapso",
    "explode", "dispara", "histórico", "historico", "inédito", "inedito",
    "polêmica", "polemica", "revela", "revelado", "expõe", "expoe", "exposto",
    "salto", "disparada", "abusiva", "abusivo", "esconde", "por dentro",
    "manobra", "brecha", "domina",
]

# Critério 3 — Rastreabilidade da Prova: documento/operação/autoridade
PROVA = [
    "operação", "operacao", "cade", "tcu", "anatel", "anvisa", "processo",
    "multa", "auditoria", "relatório", "relatorio", "investigação", "investigacao",
    "ministério público", "ministerio publico", "mpf", "justiça", "justica",
    "stf", "decisão judicial", "sentença", "sentenca", "denúncia", "denuncia",
    "polícia federal", "policia federal", "pf", "inquérito", "inquerito",
    "censo", "levantamento", "dados oficiais", "cpi", "operação policial",
]

# Critério 4 — Apelo de Massa: mexe no bolso/vida do cidadão
MASSA = [
    "preço", "preco", "bolso", "consumidor", "aluguel", "imóvel", "imovel",
    "salário", "salario", "emprego", "desemprego", "comida", "alimento",
    "energia", "combustível", "combustivel", "gasolina", "água", "agua",
    "inflação", "inflacao", "tarifa", "imposto", "renda", "família", "familia",
    "trabalhador", "população", "populacao", "consumo",
]

# >>> SINAL INVESTIGATIVO (porteiro): a "tensão" / sistema oculto <<<
# Sem ao menos UM destes, a notícia NÃO é uma pauta Elementar de verdade.
INVESTIGATIVO = [
    "investigação", "investigacao", "investiga", "escândalo", "escandalo",
    "cartel", "fraude", "esquema", "multa", "processo", "operação", "operacao",
    "abuso", "sobrepreço", "sobrepreco", "superfaturamento", "sobrecusto",
    "prejuízo", "prejuizo", "denúncia", "denuncia", "irregularidade", "suspeita",
    "conluio", "monopólio", "monopolio", "oligopólio", "oligopolio", "propina",
    "corrupção", "corrupcao", "lavagem", "sonegação", "sonegacao", "apreensão",
    "apreensao", "condenação", "condenacao", "vazamento", "manobra", "brecha",
    "atraso", "gargalo", "ineficiência", "ineficiencia", "colapso", "bolha",
    "êxodo", "exodo", "esvaziamento", "domínio", "dominio", "domina",
    "concentração", "concentracao", "falsificação", "falsificacao", "contrabando",
    "obsolescência", "obsolescencia", "lobby", "barreira", "ilegal", "golpe",
    "oculto", "secreto", "bastidores", "abusiva", "abusivo", "expõe", "expoe",
    "revela", "recall", "trava", "truste", "monopoliza", "dominar",
]

# >>> JORNALISMO DE SERVIÇO (anti-pauta): utilidade/tutorial/prazo/rotina <<<
ANTI_PAUTA = [
    "prazo", "veja como", "saiba como", "reta final", "passo a passo", "confira",
    "tira-dúvidas", "tira-duvidas", "tira dúvidas", "como declarar", "como fazer",
    "como pedir", "como solicitar", "como consultar", "calendário", "calendario",
    "último dia", "ultimo dia", "data limite", "veja o que", "dicas", "guia",
    "inscrições abertas", "inscricoes abertas", "abre inscrições",
    "horário de funcionamento", "o que abre e fecha", "feriado", "mega-sena",
    "resultado da loteria", "previsão do tempo", "previsao do tempo", "veja fotos",
    "veja vídeo", "veja video", "veja a lista", "tutorial", "como pagar",
]

# Regex p/ Critério 2 — Densidade Estatística
RE_PERCENT = re.compile(r"\d+[\.,]?\d*\s?%")
RE_DINHEIRO = re.compile(r"r\$\s?\d", re.IGNORECASE)
RE_BIMI = re.compile(r"\d+[\.,]?\d*\s?(bilh|milh|trilh)", re.IGNORECASE)
RE_NUM = re.compile(r"\d{2,}")


# ---------------------------------------------------------------------------
# 3. MOTOR DE PONTUAÇÃO  (casamento por PALAVRA INTEIRA, não pedaço)
# ---------------------------------------------------------------------------
def _compilar(termos):
    """Compila cada termo com fronteira de palavra (evita 'pf' casar em 'IRPF')."""
    return [re.compile(r"(?<!\w)" + re.escape(t) + r"(?!\w)", re.IGNORECASE)
            for t in termos]

P_GANCHO = _compilar(GANCHO)
P_PROVA = _compilar(PROVA)
P_MASSA = _compilar(MASSA)
P_INVEST = _compilar(INVESTIGATIVO)
P_ANTI = _compilar(ANTI_PAUTA)
P_CAT = {cat: _compilar(termos) for cat, termos in CATEGORIAS.items()}


def _conta(texto, patterns):
    return sum(1 for p in patterns if p.search(texto))


def _nota_0a3(qtd):
    return 0 if qtd <= 0 else min(qtd, 3)


def detectar_categoria(texto):
    melhor, melhor_qtd = "Geral", 0
    for cat, pats in P_CAT.items():
        q = _conta(texto, pats)
        if q > melhor_qtd:
            melhor, melhor_qtd = cat, q
    return melhor


def pontuar(titulo, resumo):
    titulo = titulo or ""
    resumo = resumo or ""
    full = f"{titulo} {resumo}"

    # 1. Gancho — conflito/indignação, peso maior no título
    g = _conta(titulo, P_GANCHO) * 2 + _conta(resumo, P_GANCHO)
    nota_gancho = _nota_0a3(g)

    # 2. Densidade estatística
    dens = (len(RE_PERCENT.findall(full)) + len(RE_DINHEIRO.findall(full))
            + len(RE_BIMI.findall(full)) + (1 if RE_NUM.search(full) else 0))
    nota_dens = _nota_0a3(dens)

    # 3. Rastreabilidade da prova
    nota_prova = _nota_0a3(_conta(full, P_PROVA))

    # 4. Apelo de massa
    nota_massa = _nota_0a3(_conta(full, P_MASSA))

    total = nota_gancho + nota_dens + nota_prova + nota_massa

    # >>> PORTEIRO: tensão investigativa vs. jornalismo de serviço <<<
    tensao = _conta(titulo, P_INVEST) * 2 + _conta(resumo, P_INVEST)
    servico = _conta(full, P_ANTI)
    servico_puro = (servico >= 1 and tensao == 0)

    # Recomendada só se: soma>=9 E tem conflito E não é puro serviço
    aprovada = (total >= 9) and (tensao >= 1) and (not servico_puro)

    # motivo do veredito (transparência)
    if servico_puro:
        motivo = "Jornalismo de serviço / utilidade — sem conflito investigativo."
    elif tensao == 0:
        motivo = "Sem sinal investigativo claro — números sozinhos não fazem pauta."
    elif aprovada:
        motivo = "Tem conflito + dados + autoridade. Pauta forte."
    else:
        motivo = "Tem ingredientes, mas não atinge o corte (soma ≥ 9)."

    # ranking interno: investigativo sobe, serviço afunda
    rank = total + min(tensao, 3)
    if servico_puro:
        rank -= 6

    return {
        "gancho": nota_gancho,
        "densidade": nota_dens,
        "prova": nota_prova,
        "massa": nota_massa,
        "total": total,
        "tensao": tensao,
        "servico": servico,
        "servico_puro": servico_puro,
        "aprovada": aprovada,
        "rank": rank,
        "motivo": motivo,
        "categoria": detectar_categoria(full),
    }


# ---------------------------------------------------------------------------
# 4. GERADOR DE ÂNGULO DE VÍDEO (estilo Elementar)
# ---------------------------------------------------------------------------
ANGULOS = {
    "Monopólios e Poder Econômico":
        "Gancho de contraste: o mercado quando havia concorrência vs. hoje "
        "dominado por poucos. Tese: revele a fusão/manobra que criou o domínio e "
        "como ela trava preços e barra novos entrantes.",
    "Infraestrutura e Engenharia":
        "Gancho de contraste: tempo/custo da obra no Brasil vs. no exterior. "
        "Tese: exponha o gargalo de burocracia/licitação por trás do atraso, "
        "ancorado no relatório ou auditoria como prova.",
    "Indústria do Consumo e Margens":
        "Gancho de contraste: o custo real de produção vs. o preço de prateleira. "
        "Tese: desmonte a margem inflada e o mecanismo (marca, escassez fabricada, "
        "obsolescência) que justifica o sobrepreço.",
    "Declínio Urbano e Imobiliário":
        "Gancho de contraste: a região cheia no passado vs. vazia hoje (dados do "
        "censo). Tese: explique a dinâmica econômica (êxodo, vacância, crise) e o "
        "efeito cascata no comércio e nos imóveis.",
    "Geral":
        "Gancho de contraste: parta de um número chocante da matéria. Tese: "
        "rastreie o sistema oculto (lei, decisão, esquema) por trás do fenômeno e "
        "feche com o futuro daquele mercado.",
}


def gerar_angulo(score, titulo):
    if score["servico_puro"]:
        return ("Provável pauta fraca para o canal: é notícia de utilidade/prazo, "
                "sem conflito a desvendar. Só vire vídeo se você achar um ângulo de "
                "sistema oculto por trás disso.")
    base = ANGULOS.get(score["categoria"], ANGULOS["Geral"])
    reforco = []
    if score["prova"] >= 2:
        reforco.append("Há documento/operação citável — ancore a autoridade nele.")
    if score["densidade"] >= 2:
        reforco.append("Os números já dão a abertura; comece por eles na tela.")
    if score["massa"] >= 2:
        reforco.append("Mexe no bolso do público — explore o impacto pessoal no fim.")
    if score["gancho"] == 0:
        reforco.append("Título morno: procure o contraste/indignação antes de pautar.")
    return base + ((" " + " ".join(reforco)) if reforco else "")


# ---------------------------------------------------------------------------
# 5. COLETA
# ---------------------------------------------------------------------------
def limpar(texto):
    if not texto:
        return ""
    texto = re.sub(r"<[^>]+>", " ", texto)
    texto = html.unescape(texto)
    return re.sub(r"\s+", " ", texto).strip()


def buscar_fonte(nome, urls):
    for url in urls:
        try:
            r = requests.get(url, headers=HEADERS, timeout=20)
            if r.status_code != 200:
                continue
            d = feedparser.parse(r.content)
            if d.entries:
                return d.entries, url, None
        except Exception as e:
            continue
    return [], None, "nenhum feed respondeu"


def coletar(fontes_escolhidas, limite_por_fonte=30):
    itens, status = [], {}
    for nome in fontes_escolhidas:
        entries, url_ok, erro = buscar_fonte(nome, FONTES.get(nome, []))
        status[nome] = {"ok": bool(entries), "url": url_ok,
                        "qtd": len(entries), "erro": erro}
        for e in entries[:limite_por_fonte]:
            titulo = limpar(getattr(e, "title", ""))
            resumo = limpar(getattr(e, "summary", "") or getattr(e, "description", ""))
            if not titulo:
                continue
            score = pontuar(titulo, resumo)
            itens.append({
                "fonte": nome,
                "titulo": titulo,
                "resumo": resumo[:400],
                "link": getattr(e, "link", ""),
                "data": getattr(e, "published", "") or getattr(e, "updated", ""),
                "score": score,
                "angulo": gerar_angulo(score, titulo),
            })
    itens.sort(key=lambda x: x["score"]["rank"], reverse=True)
    return itens, status


# ---------------------------------------------------------------------------
# 6. ROTAS
# ---------------------------------------------------------------------------
@app.route("/")
def index():
    return render_template("index.html", fontes=list(FONTES.keys()))


@app.route("/api/buscar", methods=["POST"])
def api_buscar():
    dados = request.get_json(force=True)
    fontes = dados.get("fontes") or list(FONTES.keys())
    nota_min = int(dados.get("nota_min", 0))
    so_aprovadas = bool(dados.get("so_aprovadas", False))
    ocultar_servico = bool(dados.get("ocultar_servico", True))

    itens, status = coletar(fontes)
    filtrados = []
    for it in itens:
        s = it["score"]
        if ocultar_servico and s["servico_puro"]:
            continue
        if so_aprovadas and not s["aprovada"]:
            continue
        if s["total"] < nota_min:
            continue
        filtrados.append(it)

    return jsonify({"itens": filtrados, "status": status, "total": len(filtrados)})


@app.route("/api/exportar", methods=["POST"])
def api_exportar():
    dados = request.get_json(force=True)
    itens = dados.get("itens", [])
    doc = Document()
    doc.add_heading("ELEMENTAR — Pautas Mineradas", level=0)
    sub = doc.add_paragraph()
    run = sub.add_run("Relatório gerado em " + datetime.now().strftime("%d/%m/%Y às %H:%M"))
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    doc.add_paragraph("Ordenadas pela força investigativa. Soma ≥ 9 + conflito = recomendada.")

    for i, it in enumerate(itens, 1):
        s = it["score"]
        doc.add_heading(f"{i}. {it['titulo']}", level=1)
        meta = doc.add_paragraph()
        mrun = meta.add_run(f"Fonte: {it['fonte']}  |  Categoria: {s['categoria']}  |  {it.get('data','')}")
        mrun.font.size = Pt(9)
        mrun.font.color.rgb = RGBColor(0x70, 0x70, 0x70)

        nota = doc.add_paragraph()
        nrun = nota.add_run(f"NOTA {s['total']}/12  ({'RECOMENDADA' if s['aprovada'] else 'fora do corte'})")
        nrun.bold = True
        nrun.font.color.rgb = (RGBColor(0x1B, 0x7A, 0x3C) if s['aprovada'] else RGBColor(0xA0, 0x40, 0x00))
        doc.add_paragraph(f"Gancho {s['gancho']}/3 · Densidade {s['densidade']}/3 · "
                          f"Rastreabilidade {s['prova']}/3 · Apelo de massa {s['massa']}/3")
        mt = doc.add_paragraph()
        mt.add_run("Veredito: ").bold = True
        mt.add_run(s["motivo"])

        if it.get("resumo"):
            doc.add_paragraph(it["resumo"])
        ang = doc.add_paragraph()
        ang.add_run("Ângulo de vídeo sugerido: ").bold = True
        ang.add_run(it.get("angulo", ""))
        if it.get("link"):
            lk = doc.add_paragraph()
            lr = lk.add_run(it["link"])
            lr.font.size = Pt(9)
            lr.font.color.rgb = RGBColor(0x1A, 0x5C, 0xA8)
        doc.add_paragraph("—" * 30)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    nome = "pautas_elementar_" + datetime.now().strftime("%Y%m%d_%H%M") + ".docx"
    return send_file(buf, as_attachment=True, download_name=nome,
                     mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


if __name__ == "__main__":
    print("\n  ELEMENTAR rodando em  http://127.0.0.1:5000\n")
    app.run(debug=True, port=5000)
